import os, io, csv, sqlite3, hmac, secrets, time, calendar, datetime as dt
from functools import wraps
from flask import (Flask, Blueprint, g, request, session, redirect, render_template,
                   jsonify, abort, flash, url_for)
from markupsafe import Markup
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import generate_password_hash, check_password_hash

ADMIN = os.environ.get("ADMIN_PATH", "").strip("/")
ADMIN_PW = os.environ.get("ADMIN_PASSWORD", "")
SECRET = os.environ.get("SECRET_KEY", "")
if len(ADMIN) < 12 or len(ADMIN_PW) < 8 or len(SECRET) < 16:
    raise SystemExit("Set ADMIN_PATH (12+ chars, secret), ADMIN_PASSWORD (8+) and SECRET_KEY (16+).")
DB = os.path.join(os.environ.get("DATA_DIR", "."), "evote.db")

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)
app.config.update(SECRET_KEY=SECRET, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=os.environ.get("INSECURE_COOKIES") != "1",
                  PERMANENT_SESSION_LIFETIME=dt.timedelta(hours=8), MAX_CONTENT_LENGTH=5 * 1024 * 1024)

SL = {"setup": "Setup", "registration": "Registration", "nomination": "Nominations",
      "voting": "Voting", "closed": "Results"}
STG = list(SL)
DEFAULT_CATS = ["President", "Vice President", "Secretary General", "Vice Secretary General",
                "Chairperson", "Marketing Manager", "Treasurer", "Project and Research Manager",
                "1st Year Representative", "Additional Member"]
SCHEMA = """
create table if not exists years(name text primary key, stage text not null default 'setup', active int not null default 0);
create table if not exists members(year text, sid text, name text, email text, campus text, level text, prog text, interests text, skill text, primary key(year,sid));
create table if not exists accounts(year text, sid text, pw text, primary key(year,sid));
create table if not exists cats(id integer primary key autoincrement, year text, name text, slots int default 3);
create table if not exists noms(id integer primary key autoincrement, year text, cat int, nominee text, t real);
create table if not exists nominated(year text, cat int, voter text, primary key(year,cat,voter));
create table if not exists cands(year text, cat int, nominee text, primary key(year,cat,nominee));
create table if not exists votes(id integer primary key autoincrement, year text, cat int, cand text);
create table if not exists voted(year text, cat int, voter text, primary key(year,cat,voter));
create table if not exists events(id integer primary key autoincrement, title text, date text, time text, place text, descr text);
"""

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB, timeout=20)
        g.db.row_factory = sqlite3.Row
        g.db.execute("pragma journal_mode=wal")
    return g.db

@app.teardown_appcontext
def _close(e):
    d = g.pop("db", None)
    if d: d.close()

with sqlite3.connect(DB) as _c:
    _c.executescript(SCHEMA)

def q1(sql, *a): return db().execute(sql, a).fetchone()
def qa(sql, *a): return db().execute(sql, a).fetchall()
def year(): return q1("select * from years where active=1")
def member(y, sid): return q1("select * from members where year=? and sid=?", y["name"], sid) if y else None

fails = {}
def throttled(k):
    now = time.time(); fails[k] = [t for t in fails.get(k, []) if now - t < 300]
    return len(fails[k]) >= 8
def fail(k): fails.setdefault(k, []).append(time.time())

@app.before_request
def guard():
    if "_csrf" not in session: session["_csrf"] = secrets.token_hex(16)
    if request.method == "POST" and not hmac.compare_digest(request.form.get("_csrf", ""), session["_csrf"]):
        abort(400)

@app.after_request
def hdr(r):
    r.headers.update({"X-Frame-Options": "DENY", "X-Content-Type-Options": "nosniff",
                      "Referrer-Policy": "same-origin", "Cache-Control": "no-store"})
    return r

@app.context_processor
def ctx():
    return dict(csrf=lambda: Markup('<input type=hidden name=_csrf value="%s">' % session["_csrf"]),
                SL=SL, STG=STG, y=year(), adm="/" + ADMIN)

def ranking(y, cat, limit=1000):
    return qa("""select n.nominee sid, m.name, count(*) n, min(n.t) t from noms n
        join members m on m.year=n.year and m.sid=n.nominee where n.year=? and n.cat=?
        group by n.nominee order by n desc, t limit ?""", y["name"], cat, limit)

def live_data(y):
    voting = y["stage"] in ("voting", "closed"); cats = []
    for c in qa("select * from cats where year=? order by id", y["name"]):
        if voting:
            rows = qa("""select m.name, (select count(*) from votes v where v.year=k.year and v.cat=k.cat and v.cand=k.nominee) n
                from cands k join members m on m.year=k.year and m.sid=k.nominee where k.year=? and k.cat=?
                order by n desc, m.name""", y["name"], c["id"])
            items = [dict(name=r["name"], n=r["n"], top=True) for r in rows]
        else:
            items = [dict(name=r["name"], n=r["n"], top=i < c["slots"]) for i, r in enumerate(ranking(y, c["id"]))]
        cats.append(dict(id=c["id"], name=c["name"], slots=c["slots"], items=items, total=sum(i["n"] for i in items)))
    n = y["name"]
    stats = [["Registered members", q1("select count(*) c from members where year=?", n)["c"]],
             ["Accounts created", q1("select count(*) c from accounts where year=?", n)["c"]],
             ["Members who voted", q1("select count(distinct voter) c from voted where year=?", n)["c"]],
             ["Votes cast", q1("select count(*) c from votes where year=?", n)["c"]]]
    return dict(stage=y["stage"], stats=stats, cats=cats)

def cal_ctx(base):
    try:
        yy, mm = map(int, (request.args.get("m") or dt.date.today().strftime("%Y-%m")).split("-")); dt.date(yy, mm, 1)
    except Exception:
        t = dt.date.today(); yy, mm = t.year, t.month
    evs = qa("select * from events where date like ? order by date, time", "%d-%02d-%%" % (yy, mm))
    by = {}
    for e in evs: by.setdefault(e["date"], []).append(e)
    p = dt.date(yy, mm, 1) - dt.timedelta(days=1); nx = dt.date(yy, mm, 28) + dt.timedelta(days=5)
    return dict(weeks=calendar.Calendar(6).monthdatescalendar(yy, mm), by=by, evs=evs, mm=mm,
                title=dt.date(yy, mm, 1).strftime("%B %Y"), today=dt.date.today(), base=base,
                prev=p.strftime("%Y-%m"), nxt=nx.strftime("%Y-%m"))

# ---------------- members ----------------
def member_required(f):
    @wraps(f)
    def w(*a, **k):
        y = year(); sid = session.get("sid")
        if not (sid and y and member(y, sid) and q1("select 1 from accounts where year=? and sid=?", y["name"], sid)):
            session.pop("sid", None); return redirect("/")
        return f(*a, **k)
    return w

@app.route("/")
def index():
    if session.get("sid"): return redirect("/dashboard")
    return render_template("auth.html")

@app.post("/login")
def login():
    y = year(); sid = request.form.get("sid", "").strip(); k = "%s|%s" % (request.remote_addr, sid)
    if throttled(k):
        flash("Too many attempts. Try again in a few minutes."); return redirect("/")
    r = y and q1("select pw from accounts where year=? and sid=?", y["name"], sid)
    if not r or not check_password_hash(r["pw"], request.form.get("pw", "")):
        fail(k); flash("Wrong Student ID or password."); return redirect("/")
    session.clear(); session["sid"] = sid; session.permanent = True
    return redirect("/dashboard")

@app.post("/register")
def register():
    y = year(); sid = request.form.get("sid", "").strip(); em = request.form.get("email", "").strip().lower()
    pw = request.form.get("pw", "")
    def no(m): flash(m); return redirect("/")
    if not y or y["stage"] not in ("registration", "nomination", "voting"): return no("Registration is not open.")
    m = member(y, sid)
    if not m: return no("Your Student ID is not on the registered members list for this academic year, so you cannot create an account.")
    if m["email"].lower() != em: return no("That email does not match your membership record.")
    if q1("select 1 from accounts where year=? and sid=?", y["name"], sid): return no("Account already exists. Please sign in.")
    if len(pw) < 6: return no("Password must be at least 6 characters.")
    with db(): db().execute("insert into accounts values(?,?,?)", (y["name"], sid, generate_password_hash(pw)))
    session.clear(); session["sid"] = sid; session.permanent = True
    return redirect("/dashboard")

@app.route("/logout")
def logout():
    session.clear(); return redirect("/")

def page(tab, **kw):
    y = year(); m = member(y, session["sid"])
    return render_template("member.html", tab=tab, user=m["name"], **kw)

@app.route("/dashboard")
@member_required
def dashboard(): return page("dashboard")

@app.route("/calendar")
@member_required
def calendar_page(): return page("calendar", cal=cal_ctx("/calendar?"))

@app.route("/nominate", methods=["GET", "POST"])
@member_required
def nominate():
    y = year()
    if request.method == "POST":
        cat = request.form.get("cat", type=int); sid = request.form.get("nominee", "")
        sid = sid[sid.rfind("(") + 1:sid.rfind(")")] if "(" in sid else sid.strip()
        if y["stage"] != "nomination" or not q1("select 1 from cats where id=? and year=?", cat, y["name"]) or not member(y, sid):
            flash("Choose a registered member from the suggestions."); return redirect("/nominate")
        try:
            with db():
                db().execute("insert into nominated values(?,?,?)", (y["name"], cat, session["sid"]))
                db().execute("insert into noms(year,cat,nominee,t) values(?,?,?,?)", (y["name"], cat, sid, time.time()))
            flash("Nomination submitted.", "ok")
        except sqlite3.IntegrityError:
            flash("You have already nominated in that category.")
        return redirect("/nominate")
    done = {r["cat"] for r in qa("select cat from nominated where year=? and voter=?", y["name"], session["sid"])}
    return page("nominate", cats=qa("select * from cats where year=? order by id", y["name"]), done=done,
                members=qa("select sid,name from members where year=? order by name", y["name"]))

@app.route("/vote", methods=["GET", "POST"])
@member_required
def vote():
    y = year()
    if request.method == "POST":
        cat = request.form.get("cat", type=int); cand = request.form.get("cand", "")
        if y["stage"] != "voting" or not q1("select 1 from cands where year=? and cat=? and nominee=?", y["name"], cat, cand):
            flash("Invalid vote."); return redirect("/vote")
        try:
            with db():  # one transaction; ballot is stored without the voter's identity
                db().execute("insert into voted values(?,?,?)", (y["name"], cat, session["sid"]))
                db().execute("insert into votes(year,cat,cand) values(?,?,?)", (y["name"], cat, cand))
            flash("Vote recorded.", "ok")
        except sqlite3.IntegrityError:
            flash("You have already voted in that category.")
        return redirect("/vote")
    done = {r["cat"] for r in qa("select cat from voted where year=? and voter=?", y["name"], session["sid"])}
    cats = []
    for c in qa("select * from cats where year=? order by id", y["name"]):
        cs = qa("select m.* from cands k join members m on m.year=k.year and m.sid=k.nominee where k.year=? and k.cat=? order by m.name", y["name"], c["id"])
        cats.append(dict(c=c, cands=cs))
    return page("vote", vcats=cats, done=done)

@app.route("/api/live")
def api_live():
    y = year()
    if not y or not (session.get("admin") or (session.get("sid") and member(y, session["sid"]))): abort(403)
    return jsonify(live_data(y))

# ---------------- admin (secret URL) ----------------
bp = Blueprint("admin", __name__, url_prefix="/" + ADMIN)

def read_rows(f):
    if f.filename.lower().endswith(".csv"):
        return list(csv.reader(io.StringIO(f.read().decode("utf-8-sig", errors="replace"))))
    from openpyxl import load_workbook
    return [list(r) for r in load_workbook(f, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)]

def parse_members(rows):
    if len(rows) < 2: raise ValueError("The file is empty.")
    hd = [str(h or "").strip().lower() for h in rows[0]]
    keys = ("student id", "full name", "email", "campus", "year of study", "programme", "areas of technology", "skill level")
    ci = {k: next((i for i, h in enumerate(hd) if k in h), None) for k in keys}
    miss = [k for k in keys[:3] if ci[k] is None]
    if miss: raise ValueError("Missing column(s): " + ", ".join(miss))
    out = {}
    for r in rows[1:]:
        def v(k):
            i = ci[k]; x = r[i] if i is not None and i < len(r) else ""
            return "" if x is None else str(x).strip()
        sid = v("student id")
        if sid.endswith(".0"): sid = sid[:-2]
        if sid and sid not in out:
            out[sid] = (sid, v("full name"), v("email").lower(), v("campus"), v("year of study"),
                        v("programme"), v("areas of technology"), v("skill level"))
    return list(out.values())

@bp.route("/", methods=["GET", "POST"])
def panel():
    if request.method == "POST" and not session.get("admin"):
        k = "admin|" + str(request.remote_addr)
        if not throttled(k) and hmac.compare_digest(request.form.get("pw", ""), ADMIN_PW):
            session.clear(); session["admin"] = True; session.permanent = True
            return redirect(url_for("admin.panel"))
        fail(k); flash("Wrong password.")
    if not session.get("admin"): return render_template("admin_login.html")
    y = year(); tab = request.args.get("tab", "dashboard"); kw = {}
    if y:
        kw["cats"] = qa("select * from cats where year=? order by id", y["name"])
        s = request.args.get("q", "").strip()
        kw["members"] = qa("select * from members where year=? and (sid||name||prog||interests) like ? order by name limit 200", y["name"], "%" + s + "%")
        kw["mcount"] = q1("select count(*) c from members where year=?", y["name"])["c"]
    kw["years"] = qa("select * from years order by name")
    if tab == "calendar": kw["cal"] = cal_ctx("%s/?tab=calendar&" % bp.url_prefix)
    return render_template("admin.html", tab=tab, q=request.args.get("q", ""), **kw)

@bp.post("/act/<a>")
def act(a):
    if not session.get("admin"): abort(404)
    y = year(); d = db(); f = request.form; tab = f.get("tab", "control")
    if a == "upload":
        name = f.get("year", "").strip(); fl = request.files.get("file")
        try:
            if not name or not fl or not fl.filename: raise ValueError("Enter the academic year and choose a file.")
            ms = parse_members(read_rows(fl))
        except Exception as e:
            flash("Upload failed: %s" % e); return redirect(url_for("admin.panel", tab="members"))
        old = [r["name"] for r in qa("select name from cats where year=? order by id", y["name"])] if y else DEFAULT_CATS
        slots = {r["name"]: r["slots"] for r in qa("select name,slots from cats where year=?", y["name"])} if y else {}
        with d:
            d.execute("update years set active=0")
            d.execute("insert into years(name,stage,active) values(?,'setup',1) on conflict(name) do update set active=1", (name,))
            d.execute("delete from members where year=?", (name,))
            d.executemany("insert into members values(?,?,?,?,?,?,?,?,?)", [(name,) + m for m in ms])
            d.execute("delete from accounts where year=? and sid not in (select sid from members where year=?)", (name, name))
            if not d.execute("select 1 from cats where year=?", (name,)).fetchone():
                d.executemany("insert into cats(year,name,slots) values(?,?,?)", [(name, c, slots.get(c, 3)) for c in old])
        flash("%d members loaded for %s." % (len(ms), name), "ok")
    elif a == "year" and f.get("name"):
        with d:
            d.execute("update years set active=0"); d.execute("update years set active=1 where name=?", (f["name"],))
    elif a == "event_add":
        if f.get("title", "").strip() and f.get("date"):
            with d: d.execute("insert into events(title,date,time,place,descr) values(?,?,?,?,?)",
                              (f["title"].strip(), f["date"], f.get("time", ""), f.get("place", ""), f.get("descr", "")))
        tab = "calendar"
    elif a == "event_del":
        with d: d.execute("delete from events where id=?", (f.get("id", type=int),))
        tab = "calendar"
    elif y:
        n = y["name"]
        if a == "stage" and f.get("stage") in STG:
            with d:
                if f["stage"] == "voting" and not d.execute("select 1 from votes where year=?", (n,)).fetchone():
                    d.execute("delete from cands where year=?", (n,))
                    for c in qa("select * from cats where year=?", n):  # most nominations first
                        for r in ranking(y, c["id"], c["slots"]):
                            d.execute("insert into cands values(?,?,?)", (n, c["id"], r["sid"]))
                d.execute("update years set stage=? where name=?", (f["stage"], n))
        elif a == "cat_add" and f.get("name", "").strip():
            with d: d.execute("insert into cats(year,name,slots) values(?,?,3)", (n, f["name"].strip()))
        elif a == "cat_slots":
            with d: d.execute("update cats set slots=? where id=? and year=?", (max(1, min(20, f.get("slots", 3, type=int))), f.get("id", type=int), n))
        elif a == "cat_del":
            cid = f.get("id", type=int)
            with d:
                for t in ("noms", "nominated", "cands", "votes", "voted"): d.execute("delete from %s where year=? and cat=?" % t, (n, cid))
                d.execute("delete from cats where id=? and year=?", (cid, n))
        elif a == "reset":
            with d:
                for t in ("accounts", "noms", "nominated", "cands", "votes", "voted"): d.execute("delete from %s where year=?" % t, (n,))
                d.execute("update years set stage='registration' where name=?", (n,))
    return redirect(url_for("admin.panel", tab=tab))

@bp.route("/logout")
def alogout():
    session.clear(); return redirect(url_for("admin.panel"))

app.register_blueprint(bp)

@app.errorhandler(404)
def nf(e): return "Not found", 404
